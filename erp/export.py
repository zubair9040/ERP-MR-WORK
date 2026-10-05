"""Excel and PDF downloads for any table, plus the statement as Excel."""
import io
import re
from datetime import datetime

from flask import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from xml.sax.saxutils import escape

from .db import fmt, nice_date, settings

MONEY_FMT = '#,##0.00;[Red]-#,##0.00'
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
THIN = Side(style="thin", color="BBBBBB")


def _filename(title, ext):
    base = re.sub(r"[^A-Za-z0-9]+", "_", title).strip("_") or "export"
    return f"{base}_{datetime.now():%Y%m%d}.{ext}"


def _send(data, mimetype, name):
    return Response(data, mimetype=mimetype, headers={"Content-Disposition": f'attachment; filename="{name}"'})


def _xl_value(v, is_money):
    if is_money and isinstance(v, int):
        return v / 100
    return v


def table_xlsx(title, headers, rows, money_cols=(), totals=None, subtitle=""):
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\\/*?:\[\]]", "", title)[:31] or "Sheet1"
    ws.append([settings().get("company_name", "")])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([title])
    ws["A2"].font = Font(bold=True, size=12)
    ws.append([subtitle or f"Printed {datetime.now():%d-%b-%Y %H:%M}"])
    ws.append([])
    ws.append(list(headers))
    hr = ws.max_row
    for k in range(1, len(headers) + 1):
        c = ws.cell(row=hr, column=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append([_xl_value(v, i in money_cols) for i, v in enumerate(r)])
    if totals:
        ws.append([_xl_value(v, i in money_cols) for i, v in enumerate(totals)])
        for k in range(1, len(headers) + 1):
            c = ws.cell(row=ws.max_row, column=k)
            c.font = Font(bold=True)
            c.border = Border(top=Side(style="medium"))
    for i in money_cols:
        for row in ws.iter_rows(min_row=hr + 1, min_col=i + 1, max_col=i + 1):
            for c in row:
                c.number_format = MONEY_FMT
    for k, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(f"{r[k - 1]:,.2f}" if (k - 1) in money_cols and isinstance(r[k - 1], (int, float))
                                          else str(r[k - 1])) for r in rows[:500]] or [10])
        ws.column_dimensions[get_column_letter(k)].width = min(max(width + 2, 9), 60)
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def table_pdf(title, headers, rows, money_cols=(), totals=None, subtitle=""):
    from .companies import pdf_settings
    from .pdfs import GREY, INK, LINE, SOFT, _accent, _f, _image
    s = pdf_settings()
    ac = _accent(s)
    wide = len(headers) > 5
    page = landscape(A4) if wide else A4
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm,
                            bottomMargin=12 * mm, title=title)
    cell = ParagraphStyle("c", fontName=_f(), fontSize=7.5, leading=9.5, textColor=INK)
    head = ParagraphStyle("h", parent=cell, fontName=_f(True), textColor=colors.white)

    def txt(i, v):
        if i in money_cols and isinstance(v, int):
            return fmt(v)
        return Paragraph(escape(str(v if v is not None else "")).replace("\n", "<br/>"), cell)

    data = [[Paragraph(escape(str(h)), head) for h in headers]]
    data += [[txt(i, v) for i, v in enumerate(r)] for r in rows]
    if totals:
        data.append([Paragraph(escape(fmt(v) if i in money_cols and isinstance(v, int) else str(v)),
                               ParagraphStyle("t", parent=cell, fontName=_f(True), alignment=2 if i in money_cols else 0))
                     for i, v in enumerate(totals)])
    avail = page[0] - 24 * mm
    lens = [max([len(str(h))] + [len(str(r[i])) for r in rows[:200]] or [6]) for i, h in enumerate(headers)]
    lens = [min(max(l, 6), 45) for l in lens]
    widths = [avail * l / sum(lens) for l in lens]
    t = Table(data, colWidths=widths, repeatRows=1)
    st = [("BACKGROUND", (0, 0), (-1, 0), ac), ("FONTSIZE", (0, 0), (-1, -1), 7.5),
          ("LINEBELOW", (0, 1), (-1, -1), 0.4, LINE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
          ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
          ("ROWBACKGROUNDS", (0, 1), (-1, -1 if not totals else -2), [colors.white, colors.HexColor("#fafbfd")])]
    for i in money_cols:
        st.append(("ALIGN", (i, 1), (i, -1), "RIGHT"))
    if totals:
        st += [("BACKGROUND", (0, -1), (-1, -1), SOFT), ("LINEABOVE", (0, -1), (-1, -1), 1, INK)]
    t.setStyle(TableStyle(st))
    co = [Paragraph(escape(s.get("company_name", "")), ParagraphStyle("co", fontName=_f(True), fontSize=12, leading=15, textColor=INK)),
          Paragraph(escape(" · ".join(x for x in (s.get("company_address", "").replace("\n", ", "), s.get("company_phone")) if x)),
                    ParagraphStyle("ca", fontName=_f(), fontSize=7.5, leading=9.5, textColor=GREY))]
    right = [Paragraph(escape(title.upper()), ParagraphStyle("ti", fontName=_f(True), fontSize=13, leading=16, textColor=ac, alignment=2)),
             Paragraph(escape(subtitle or f"Printed {datetime.now():%d-%b-%Y %H:%M}"),
                       ParagraphStyle("su", fontName=_f(), fontSize=7.5, leading=9.5, textColor=GREY, alignment=2))]
    logo = _image(s["_logo"], 32 * mm, 16 * mm) if s.get("_logo") else None
    avail_w = page[0] - 24 * mm
    if logo:
        logo.hAlign = "LEFT"
        head_t = Table([[logo, co, right]], colWidths=[36 * mm, avail_w * 0.5 - 36 * mm, avail_w * 0.5])
    else:
        head_t = Table([[co, right]], colWidths=[avail_w * 0.5, avail_w * 0.5])
    head_t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("LINEBELOW", (0, 0), (-1, 0), 0.6, LINE),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    story = [head_t, Spacer(1, 8), t]

    def page_no(canvas, d):
        canvas.saveState()
        canvas.setFillColor(ac)
        canvas.rect(0, page[1] - 2.5 * mm, page[0], 2.5 * mm, stroke=0, fill=1)
        canvas.setFont(_f(), 7)
        canvas.setFillColor(GREY)
        canvas.drawString(12 * mm, 7 * mm, s.get("company_name", ""))
        canvas.drawRightString(page[0] - 12 * mm, 7 * mm, f"Page {d.page}")
        canvas.restoreState()
    doc.build(story, onFirstPage=page_no, onLaterPages=page_no)
    return buf.getvalue()


def table_response(fmt_, title, headers, rows, money_cols=(), totals=None, subtitle=""):
    if fmt_ == "pdf":
        return _send(table_pdf(title, headers, rows, money_cols, totals, subtitle), "application/pdf", _filename(title, "pdf"))
    return _send(table_xlsx(title, headers, rows, money_cols, totals, subtitle),
                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _filename(title, "xlsx"))


def statement_xlsx(customer, st, aging=None, details=None):
    from .ledger import BUCKETS
    from .lines import price_text, qty_text
    from .pdfs import _tx_text
    s = settings()
    cur = s.get("currency", "Rs")
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    bold = Font(bold=True)
    ws.append([s.get("company_name", "")])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["STATEMENT", "", "", f"Date: {nice_date(st['end'])}"])
    ws["A2"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["To:", customer["name"]])
    ws["B4"].font = bold
    for ln in (customer["address"] or "").splitlines():
        ws.append(["", ln])
    if customer["whatsapp"] or customer["phone"]:
        ws.append(["", customer["whatsapp"] or customer["phone"]])
    ws.append([])
    ws.append(["Amount due", st["closing"] / 100])
    ws.cell(row=ws.max_row, column=2).number_format = MONEY_FMT
    ws.cell(row=ws.max_row, column=1).font = bold
    ws.cell(row=ws.max_row, column=2).font = bold
    ws.append([f"All open transactions as of {nice_date(st['end'])}" if st.get("open_mode")
               else f"Statement period {nice_date(st['start'])} to {nice_date(st['end'])}"])
    ws.append([])
    ws.append(["Date", "Transaction", "Amount", "Balance"] + (["Item amount"] if details else []))
    hr = ws.max_row
    for k in range(1, 6 if details else 5):
        c = ws.cell(row=hr, column=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
    if not st.get("open_mode"):
        ws.append([nice_date(st["start"]), "Balance forward", None, st["brought_forward"] / 100])
    for l in st["lines"]:
        ws.append([nice_date(l["date"]), _tx_text(l), (l["debit"] - l["credit"]) / 100, l["balance"] / 100])
        if details and l["type"] == "Invoice" and l.get("id") in details:
            for d in details[l["id"]]:
                t = " ".join(x for x in [qty_text(d), d["unit"] or "", d["code"] or "", "-", d["description"] or ""] if x)
                pr = price_text(d)
                ws.append(["", f"    {t}" + (f" @ {pr}" if pr and d["kind"] != "subtotal" else ""), None, None, d["amount"] / 100])
                ws.cell(row=ws.max_row, column=2).font = Font(italic=True, color="555555")
                ws.cell(row=ws.max_row, column=5).number_format = MONEY_FMT
    for row in ws.iter_rows(min_row=hr + 1, min_col=3, max_col=4):
        for c in row:
            c.number_format = MONEY_FMT
    if aging:
        ws.append([])
        keys = list(aging.keys())
        labels = ([f"{k} days old" if k != "180+" else "Over 180 days" for k in keys] if "0-30" in aging else
                  ["Current", "1-30 days past due", "31-60 days past due", "61-90 days past due", "Over 90 days past due"]) + ["Amount due"]
        ws.append(labels)
        for k in range(1, 7):
            ws.cell(row=ws.max_row, column=k).font = bold
            ws.cell(row=ws.max_row, column=k).alignment = Alignment(wrap_text=True, horizontal="center")
        ws.append([aging[b] / 100 for b in keys] + [st["closing"] / 100])
        for k in range(1, 7):
            ws.cell(row=ws.max_row, column=k).number_format = MONEY_FMT
    for col_, w in zip("ABCDEF", (14, 62, 16, 18, 16, 16)):
        ws.column_dimensions[col_].width = w
    buf = io.BytesIO()
    wb.save(buf)
    name = _filename(f"Statement {customer['name']}", "xlsx")
    return _send(buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", name)
