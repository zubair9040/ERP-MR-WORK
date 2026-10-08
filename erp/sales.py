"""Customers, items, invoices and payments."""
from datetime import date, timedelta

from flask import (Blueprint, abort, flash, g, jsonify, redirect, render_template, request,
                   send_file, url_for)

from . import notify
from .auth import can, check_customer_access, is_rep, rep_filter, roles
from .db import commit, fmt, next_number, pct_rnd, nice_date, now, parse_date, plain, q, settings, to_paisa, today, x
from .ledger import apply_payments, balance, customers_with_balance, open_items, statement
from .export import table_response
from .companies import co_images, co_settings, companies, customer_company, default_company_id, get_company, next_invoice_number
from .lines import compute, parse_price, parse_qty, price_text, qty_text

bp = Blueprint("sales", __name__)

METHODS = ["Cash", "Cheque", "Bank transfer", "Online", "Other"]


def _fmt():
    """PDF or picture chosen on the spot (blank = the Settings default)."""
    v = request.form.get("fmt")
    return v if v in ("pdf", "image") else None


def staff_only():
    if is_rep():
        abort(403)


def get_customer(cid):
    return check_customer_access(q("SELECT * FROM customers WHERE id = ?", (cid,), one=True))


def scoped_customer_list(active_only=True):
    where, args = rep_filter()
    if active_only:
        where += " AND c.active = 1"
    return q(f"""SELECT c.id, c.name, c.rep_id, c.terms_days, c.whatsapp, c.company_id FROM customers c
                 WHERE 1 = 1 {where} ORDER BY c.name COLLATE NOCASE""", args)


def inv_status(inv, open_amt):
    if inv["void"]:
        return "void", "Void"
    if open_amt <= 0:
        return "paid", "Paid"
    if (inv["due_date"] or inv["date"]) < today():
        return "over", "Overdue"
    if open_amt < inv["total"]:
        return "open", "Part paid"
    return "open", "Open"


# ---- customers ---------------------------------------------------------------------------
@bp.route("/customers")
def customers():
    where, args = rep_filter()
    a = request.args
    s = a.get("q", "").strip()
    if s:
        where += " AND (c.name LIKE :s OR c.code LIKE :s OR c.whatsapp LIKE :s OR c.phone LIKE :s OR c.contact LIKE :s OR c.address LIKE :s)"
        args["s"] = f"%{s}%"
    reps_sel = [int(v) for v in a.getlist("rep") if v.isdigit()]
    if reps_sel and not is_rep():
        where += f" AND c.rep_id IN ({','.join(str(v) for v in reps_sel)})"
    cities = [v for v in a.getlist("city") if v]
    if cities:
        marks = []
        for n, v in enumerate(cities):
            args[f"city{n}"] = v
            marks.append(f":city{n}")
        where += f" AND c.city IN ({','.join(marks)})"
    if not a.get("inactive"):
        where += " AND c.active = 1"
    sort = a.get("sort", "name")
    order = {"name": "c.name COLLATE NOCASE", "name_desc": "c.name COLLATE NOCASE DESC", "code": "c.code",
             "city": "c.city COLLATE NOCASE, c.name COLLATE NOCASE"}.get(sort, "c.name COLLATE NOCASE")
    rows = customers_with_balance(None, where, args, order=order)
    if a.get("due"):
        rows = [r for r in rows if r["balance"] > 0]
    if sort in ("bal_desc", "bal_asc"):
        rows = sorted(rows, key=lambda r: r["balance"], reverse=sort == "bal_desc")
    if a.get("export") in ("xlsx", "pdf"):
        data = [[r["code"] or "", r["name"], r["city"] or "", r["contact"] or "", r["whatsapp"] or "", r["phone"] or "", r["rep_code"] or "",
                 r["balance"]] for r in rows]
        return table_response(a["export"], "Customers", ["Code", "Customer", "City", "Contact", "WhatsApp", "Phone", "Rep", "Balance"],
                              data, {7}, ["Total", "", "", "", "", "", "", sum(r["balance"] for r in rows)])
    reps = q("SELECT * FROM reps ORDER BY name")
    all_cities = [r["city"] for r in q("SELECT DISTINCT city FROM customers WHERE city != '' ORDER BY city COLLATE NOCASE")]
    filtered = bool(s or reps_sel or cities or a.get("due") or a.get("inactive") or a.get("sort", "name") != "name")
    return render_template("customers.html", rows=rows, reps=reps, total=sum(r["balance"] for r in rows), all_cities=all_cities,
                           reps_sel=reps_sel, cities_sel=cities, sort=sort, filtered=filtered)


@bp.route("/customers/new", methods=["GET", "POST"])
@bp.route("/customers/<int:cid>/edit", methods=["GET", "POST"])
def customer_form(cid=None):
    staff_only()
    c = get_customer(cid) if cid else None
    reps = q("SELECT * FROM reps WHERE active = 1 OR id = ? ORDER BY name", (c["rep_id"] if c else -1,))
    if request.method == "POST":
        f = request.form
        try:
            vals = dict(
                name=f["name"].strip(), code=f.get("code", "").strip().upper(), city=f.get("city", "").strip(),
                contact=f.get("contact", "").strip(), whatsapp=f.get("whatsapp", "").strip(),
                phone=f.get("phone", "").strip(), email=f.get("email", "").strip(), address=f.get("address", "").strip(),
                rep_id=int(f["rep_id"]) if f.get("rep_id") else None, terms_days=int(f.get("terms_days") or 0),
                credit_limit=to_paisa(f.get("credit_limit")), opening_balance=to_paisa(f.get("opening_balance")),
                opening_date=parse_date(f.get("opening_date")), notes=f.get("notes", "").strip(),
                active=1 if f.get("active") else 0,
                payment_mode="billwise" if f.get("payment_mode") == "billwise" else "fifo",
                statement_mode="open" if f.get("statement_mode") == "open" else "activity",
                send_by=f.get("send_by") if f.get("send_by") in ("whatsapp", "email", "both") else "whatsapp",
                company_id=f.get("company_id", type=int) or (c["company_id"] if c else default_company_id()),
            )
            if not get_company(vals["company_id"]):
                raise ValueError("Please choose the company this customer buys from.")
            if c and c["company_id"] and vals["company_id"] != c["company_id"] and q(
                    "SELECT 1 FROM invoices WHERE customer_id = ? UNION SELECT 1 FROM payments WHERE customer_id = ?", (cid, cid)):
                raise ValueError("This customer already has invoices or payments under another company, so the company can't be changed. "
                                 "Make a new customer under the other company instead.")
        except ValueError as e:
            flash(str(e), "err")
            return render_template("customer_form.html", c=c, reps=reps, f=f)
        from .db import guess_city, next_customer_code
        if not vals["code"]:
            vals["code"] = c["code"] if c and c["code"] else next_customer_code()
        if not vals["city"]:
            vals["city"] = guess_city(vals["address"])
        if q("SELECT 1 FROM customers WHERE code = ? COLLATE NOCASE AND id != ?", (vals["code"], cid or -1)):
            flash(f"Customer code {vals['code']} is already used by another customer.", "err")
            return render_template("customer_form.html", c=c, reps=reps, f=f)
        if not vals["name"]:
            flash("Name is required.", "err")
            return render_template("customer_form.html", c=c, reps=reps, f=f)
        if vals["email"] and not notify.valid_emails(vals["email"]):
            flash("That email address doesn't look right. You can put more than one, separated by commas.", "err")
            return render_template("customer_form.html", c=c, reps=reps, f=f)
        if vals["send_by"] != "whatsapp" and not vals["email"]:
            flash("Add an email address to send this customer's documents by email.", "err")
            return render_template("customer_form.html", c=c, reps=reps, f=f)
        if vals["whatsapp"] and not notify.normalize(vals["whatsapp"]):
            flash("That WhatsApp number doesn't look right. Use e.g. 0300-1234567 or +92 300 1234567.", "err")
            return render_template("customer_form.html", c=c, reps=reps, f=f)
        cols = ", ".join(f"{k} = :{k}" for k in vals)
        if c:
            x(f"UPDATE customers SET {cols} WHERE id = :id", {**vals, "id": cid})
        else:
            cid = x(f"INSERT INTO customers ({', '.join(vals)}, created_at) VALUES ({', '.join(':' + k for k in vals)}, :ca)",
                    {**vals, "ca": now()})
        commit()
        flash("Customer saved.", "ok")
        return redirect(url_for("sales.customer_view", cid=cid))
    return render_template("customer_form.html", c=c, reps=reps, f={})


@bp.route("/customers/<int:cid>")
def customer_view(cid):
    c = get_customer(cid)
    rep = q("SELECT * FROM reps WHERE id = ?", (c["rep_id"],), one=True) if c["rep_id"] else None
    tab = request.args.get("tab", "statement")
    t = date.today()
    start = parse_date(request.args.get("from")) or (t.replace(day=1) - timedelta(days=60)).replace(day=1).isoformat()
    end = parse_date(request.args.get("to")) or t.isoformat()
    open_by_inv, opening_left, credit = open_items(cid)
    invs = q("SELECT * FROM invoices WHERE customer_id = ? ORDER BY date DESC, number DESC", (cid,))
    inv_rows = []
    for i in invs:
        o = open_by_inv.get(i["id"], 0)
        cls, label = inv_status(i, o)
        inv_rows.append({"i": i, "open": o, "cls": cls, "label": label, "sent": notify.last_sent("invoice", i["id"])})
    pays = q("SELECT * FROM payments WHERE customer_id = ? ORDER BY date DESC, number DESC", (cid,))
    branches = q("SELECT * FROM customer_branches WHERE customer_id = ? ORDER BY active DESC, name COLLATE NOCASE", (cid,))
    pos = q("SELECT * FROM po_batches WHERE customer_id = ? AND void = 0 ORDER BY date DESC, id DESC", (cid,))
    return render_template("customer_view.html", c=c, rep=rep, tab=tab, st=statement(cid, start, end), branches=branches, pos=pos,
                           bal=balance(cid), inv_rows=inv_rows, pays=pays, opening_left=opening_left, credit=credit,
                           wa=notify.normalize(c["whatsapp"]) or "", presets=_period_presets(t))


def _period_presets(t):
    this_m = t.replace(day=1)
    last_m_end = this_m - timedelta(days=1)
    last_m = last_m_end.replace(day=1)
    return [("This month", this_m.isoformat(), t.isoformat()),
            ("Last month", last_m.isoformat(), last_m_end.isoformat()),
            ("Last 3 months", (this_m - timedelta(days=62)).replace(day=1).isoformat(), t.isoformat()),
            ("This year", t.replace(month=1, day=1).isoformat(), t.isoformat())]


def statement_options(src):
    """Reads the statement option boxes from a form or query string."""
    start = parse_date(src.get("from")) or date.today().replace(day=1).isoformat()
    end = parse_date(src.get("to")) or today()
    mode = src.get("mode") if src.get("mode") in ("period", "open", "balance") else "period"
    aging = src.getlist("aging") if hasattr(src, "getlist") else [src.get("aging")]
    opts = {"aging": (aging[-1] if aging and aging[-1] is not None else "1") == "1", "details": src.get("details") == "1"}
    return start, end, mode, opts


@bp.route("/customers/<int:cid>/statement.pdf")
def statement_pdf(cid):
    get_customer(cid)
    start, end, mode, opts = statement_options(request.args)
    return send_file(notify.statement_pdf(cid, start, end, "open" if mode == "open" else "period", opts),
                     mimetype="application/pdf")


@bp.route("/customers/<int:cid>/statement.xlsx")
def statement_xlsx(cid):
    get_customer(cid)
    start, end, mode, opts = statement_options(request.args)
    from .export import statement_xlsx as build
    c, st, aging, details = notify.statement_data(cid, start, end, "open" if mode == "open" else "period", opts)
    return build(c, st, aging, details)


@bp.route("/customers/<int:cid>/statement/send", methods=["POST"])
def statement_send(cid):
    get_customer(cid)
    start, end, mode, opts = statement_options(request.form)
    if request.form.get("channel") == "email":
        if mode == "balance":
            ok, msg = notify.email_balance(cid, request.form.get("email"))
        else:
            ok, msg = notify.email_statement(cid, start, end, request.form.get("email"), mode, opts)
    elif mode == "balance":
        ok, msg = notify.send_balance(cid, request.form.get("phone"))
    else:
        ok, msg = notify.send_statement(cid, start, end, request.form.get("phone"), mode, opts)
    flash(msg, "ok" if ok else "err")
    return redirect(url_for("sales.customer_view", cid=cid, **{"from": start, "to": end}))


@bp.route("/customers/<int:cid>/whatsapp/<what>", methods=["POST"])
def customer_whatsapp(cid, what):
    get_customer(cid)
    if what == "balance":
        ok, msg = notify.send_balance(cid, request.form.get("phone"))
    elif what == "statement":
        ok, msg = notify.send_quick_statement(cid, request.form.get("phone"))
    elif what == "email_balance":
        ok, msg = notify.email_balance(cid)
    elif what == "email_statement":
        c_ = get_customer(cid)
        start, end, mode = notify.default_period(c_)
        ok, msg = notify.email_statement(cid, start, end, None, mode)
    else:
        abort(404)
    flash(msg, "ok" if ok else "err")
    back = request.form.get("back", "")
    return redirect(back if back.startswith("/") and not back.startswith("//") else url_for("sales.customer_view", cid=cid))


@bp.route("/api/customer/<int:cid>")
def customer_api(cid):
    c = get_customer(cid)
    open_by_inv, opening_left, credit = open_items(cid, exclude_payment=request.args.get("exclude", type=int))
    open_list = []
    ids = [k for k, v in open_by_inv.items() if v]
    if ids:
        for i in q(f"""SELECT i.id, i.number, i.date, i.due_date, i.total, b.id AS po_id, b.po_no, b.combined_no
                        FROM invoices i LEFT JOIN po_batches b ON b.id = i.batch_id AND b.void = 0
                        WHERE i.id IN ({','.join('?' * len(ids))}) ORDER BY i.date, i.number""", ids):
            open_list.append({"id": i["id"], "number": i["number"], "date": nice_date(i["date"]),
                              "due": nice_date(i["due_date"]), "overdue": (i["due_date"] or i["date"]) < today(),
                              "total": plain(i["total"]), "open": plain(open_by_inv[i["id"]]),
                              "po": ({"id": i["po_id"], "po_no": i["po_no"], "combined_no": i["combined_no"]}
                                     if i["po_id"] and i["combined_no"] else None)})
    recent = q("""SELECT * FROM (
                    SELECT 'Invoice' AS type, id, number, date, total AS amount FROM invoices WHERE customer_id = :c AND void = 0
                    UNION ALL
                    SELECT 'Payment', id, number, date, amount FROM payments WHERE customer_id = :c AND void = 0)
                  ORDER BY date DESC, number DESC LIMIT 6""", {"c": cid})
    co = get_company(c["company_id"])
    return jsonify(company=({"id": co["id"], "name": co["name"], "code": co["code"], "color": co["color"], "gst": co["gst_registered"],
                             "tax": co["default_tax_rate"], "next_invoice": next_invoice_number(co["id"])} if co else None),
                   id=c["id"], name=c["name"], payment_mode=c["payment_mode"] or "fifo", address=c["address"], phone=c["whatsapp"] or c["phone"],
                   rep_id=c["rep_id"], terms_days=c["terms_days"], whatsapp=c["whatsapp"],
                   balance=plain(balance(cid)), credit_limit=plain(c["credit_limit"]),
                   opening_open=plain(opening_left), credit=plain(credit), open_invoices=open_list,
                   branches=[{"id": b["id"], "name": b["name"], "code": b["code"], "address": b["address"], "phone": b["phone"]}
                             for b in q("SELECT * FROM customer_branches WHERE customer_id = ? AND active = 1 ORDER BY name", (cid,))],
                   recent=[{"type": r["type"], "date": nice_date(r["date"]), "number": r["number"],
                            "amount": plain(r["amount"]),
                            "url": url_for("sales.invoice_view" if r["type"] == "Invoice" else "sales.payment_view",
                                           **({"iid": r["id"]} if r["type"] == "Invoice" else {"pid": r["id"]}))}
                           for r in recent])


# ---- items -------------------------------------------------------------------------------
@bp.route("/items")
def items():
    a = request.args
    s = a.get("q", "").strip()
    sql, args = "SELECT * FROM items WHERE 1 = 1", []
    if s:
        sql += " AND (name LIKE ? OR code LIKE ? OR description LIKE ? OR grp LIKE ?)"
        args += [f"%{s}%"] * 4
    groups_sel = [g_ for g_ in a.getlist("group") if g_]
    if groups_sel:
        sql += f" AND COALESCE(grp, '') IN ({','.join('?' * len(groups_sel))})"
        args += [("" if g_ == "(none)" else g_) for g_ in groups_sel]
    if not a.get("inactive"):
        sql += " AND active = 1"
    order = "COALESCE(NULLIF(grp, ''), 'zzz') COLLATE NOCASE, code" if a.get("sort") == "group" else "kind != 'item', code, name COLLATE NOCASE"
    return render_template("items.html", rows=q(sql + " ORDER BY " + order, args), kinds=ITEM_KINDS, groups=item_groups(),
                           groups_sel=groups_sel, filtered=bool(s or groups_sel or a.get("inactive") or a.get("sort")))


def item_groups():
    return [r["grp"] for r in q("SELECT DISTINCT grp FROM items WHERE grp != '' ORDER BY grp COLLATE NOCASE")]


@bp.route("/items/set-group", methods=["POST"])
def items_set_group():
    """Put the ticked items into a group (or take them out with an empty group)."""
    ids = [int(v) for v in request.form.getlist("iid") if v.isdigit()]
    grp = " ".join(request.form.get("grp", "").split())[:60]
    if not ids:
        flash("Tick the items first.", "err")
    else:
        x(f"UPDATE items SET grp = ? WHERE id IN ({','.join('?' * len(ids))})", [grp] + ids)
        commit()
        flash(f"{len(ids)} item(s) " + (f"put in group '{grp}'." if grp else "taken out of their group."), "ok")
    back = request.form.get("back", "")
    return redirect(back if back.startswith("/") and not back.startswith("//") else url_for("sales.items"))


ITEM_KINDS = {"item": "Item / product", "service": "Service or charge", "subtotal": "Sub total line",
              "discount": "Discount (percent)"}


@bp.route("/items/new", methods=["GET", "POST"])
@bp.route("/items/<int:item_id>/edit", methods=["GET", "POST"])
def item_form(item_id=None):
    staff_only()
    it = q("SELECT * FROM items WHERE id = ?", (item_id,), one=True) if item_id else None
    if item_id and not it:
        abort(404)
    if request.method == "POST":
        f = request.form
        kind = f.get("kind") if f.get("kind") in ITEM_KINDS else "item"
        try:
            percent = float(f.get("percent") or 0) if kind == "discount" else 0
            if kind == "discount" and percent > 0:
                percent = -percent
            vals = (f.get("code", "").strip().upper(), f["name"].strip(), f.get("description", "").strip(),
                    f.get("unit", "").strip(), to_paisa(f.get("price")) if kind in ("item", "service") else 0,
                    kind, percent, 1 if f.get("active") else 0,
                    float(f.get("reorder_level") or 0), to_paisa(f.get("opening_cost")),
                    f.get("barcode", "").strip(), f.get("location", "").strip(), float(f.get("pcs_per_ctn") or 0),
                    " ".join(f.get("grp", "").split())[:60])
        except ValueError as e:
            flash(str(e) if "valid" in str(e) else "Percent, reorder level and opening quantity must be numbers.", "err")
            return render_template("item_form.html", it=it, kinds=ITEM_KINDS, groups=item_groups())
        dup = q("SELECT id FROM items WHERE code = ? AND code != '' AND id != ?", (vals[0], item_id or -1), one=True)
        if vals[-4] and q("SELECT 1 FROM items WHERE barcode = ? AND id != ?", (vals[-4], item_id or -1)):
            flash(f"Barcode {vals[-4]} is already used by another item.", "err")
            return render_template("item_form.html", it=it, kinds=ITEM_KINDS, groups=item_groups())
        if dup:
            flash(f"Item code {vals[0]} is already used by another item.", "err")
            return render_template("item_form.html", it=it, kinds=ITEM_KINDS, groups=item_groups())
        if it:
            x("""UPDATE items SET code=?, name=?, description=?, unit=?, price=?, kind=?, percent=?, active=?, reorder_level=?,
                 opening_cost=?, barcode=?, location=?, pcs_per_ctn=?, grp=? WHERE id=?""", vals + (item_id,))
        else:
            x("""INSERT INTO items (code, name, description, unit, price, kind, percent, active, reorder_level,
                 opening_cost, barcode, location, pcs_per_ctn, grp) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", vals)
        commit()
        flash("Item saved.", "ok")
        return redirect(url_for("sales.items"))
    return render_template("item_form.html", it=it, kinds=ITEM_KINDS, groups=item_groups())


# ---- invoices ----------------------------------------------------------------------------
@bp.route("/invoices")
def invoices():
    where, args = rep_filter()
    a = request.args
    if a.get("q"):
        where += " AND (c.name LIKE :s OR CAST(i.number AS TEXT) = :n)"
        args.update(s=f"%{a['q'].strip()}%", n=a["q"].strip())
    if a.get("rep", type=int) and not is_rep():
        where += " AND i.rep_id = :rep"
        args["rep"] = a.get("rep", type=int)
    if a.get("delivered") == "yes":
        where += " AND i.delivered_on != ''"
    elif a.get("delivered") == "no":
        where += " AND i.delivered_on = '' AND i.void = 0"
    if parse_date(a.get("from")):
        where += " AND i.date >= :from"
        args["from"] = a["from"]
    if parse_date(a.get("to")):
        where += " AND i.date <= :to"
        args["to"] = a["to"]
    rows = q(f"""SELECT i.*, c.name AS customer, r.code AS rep_code FROM invoices i
                 JOIN customers c ON c.id = i.customer_id LEFT JOIN reps r ON r.id = i.rep_id
                 WHERE 1 = 1 {where} ORDER BY i.date DESC, i.number DESC LIMIT 500""", args)
    cache, out = {}, []
    status = a.get("status", "")
    for i in rows:
        if i["customer_id"] not in cache:
            cache[i["customer_id"]] = open_items(i["customer_id"])[0]
        o = cache[i["customer_id"]].get(i["id"], 0)
        cls, label = inv_status(i, o)
        if status and not ((status == "open" and cls in ("open", "over")) or status == cls):
            continue
        out.append({"i": i, "open": 0 if i["void"] else o, "cls": cls, "label": label})
    if a.get("export") in ("xlsx", "pdf"):
        data = [[r["i"]["number"], nice_date(r["i"]["date"]), r["i"]["customer"], r["i"]["rep_code"] or "",
                 nice_date(r["i"]["due_date"]), r["i"]["total"], r["open"], r["label"],
                 nice_date(r["i"]["delivered_on"]) if r["i"]["delivered_on"] else "Not delivered"] for r in out]
        return table_response(a["export"], "Invoices", ["Num", "Date", "Customer", "Rep", "Due", "Total", "Open", "Status", "Delivered"],
                              data, {5, 6}, ["Total", "", "", "", "", sum(r["i"]["total"] for r in out if not r["i"]["void"]),
                                             sum(r["open"] for r in out), "", ""])
    return render_template("invoices.html", rows=out, reps=q("SELECT * FROM reps ORDER BY name"),
                           total=sum(r["i"]["total"] for r in out if not r["i"]["void"]),
                           open_total=sum(r["open"] for r in out))


LINE_FIELDS = ("qty", "unit", "ctn_qty", "code", "description", "price", "item_id", "kind")


def _form_lines(f):
    """Raw line values from the posted grid, for re-showing the form."""
    cols = [f.getlist(k) for k in LINE_FIELDS]
    n = max((len(c) for c in cols), default=0)
    out = []
    for i in range(n):
        row = {k: (cols[j][i] if i < len(cols[j]) else "") for j, k in enumerate(LINE_FIELDS)}
        if any(str(row[k]).strip() for k in ("qty", "code", "description", "price", "ctn_qty")):
            out.append(row)
    return out


def _read_lines(f):
    items_by_code = {r["code"].upper(): r for r in q("SELECT * FROM items WHERE code != ''")}
    lines = []
    for n, row in enumerate(_form_lines(f), start=1):
        code = row["code"].strip().upper()
        it = items_by_code.get(code) if code else None
        kind = it["kind"] if it else "item"
        try:
            qty = parse_qty(row["qty"])
            rate, percent = parse_price(row["price"])
        except ValueError as e:
            raise ValueError(f"Line {n}: {e}")
        if kind == "subtotal":
            qty, rate, percent = None, 0, None
        elif kind == "discount" and percent is None and not rate:
            percent = it["percent"]
        lines.append({"item_id": it["id"] if it else None, "code": code, "description": row["description"].strip(),
                      "unit": row["unit"].strip(), "ctn_qty": row["ctn_qty"].strip(), "qty": qty, "rate": rate,
                      "percent": percent, "kind": "subtotal" if kind == "subtotal" else (
                          "discount" if percent is not None else kind), "sort": n})
    compute(lines)
    return lines


def _grid_rows(iid):
    rows = []
    for l in q("SELECT * FROM invoice_lines WHERE invoice_id = ? ORDER BY sort", (iid,)):
        rows.append({"qty": qty_text(l), "unit": l["unit"] or "", "ctn_qty": l["ctn_qty"] or "",
                     "code": l["code"] or "", "description": l["description"] or "", "price": price_text(l),
                     "item_id": l["item_id"] or "", "kind": l["kind"] or "item"})
    return rows


def _nav_ids(iid):
    """Previous / next invoice (by number) that this user may see."""
    where, args = rep_filter()
    cur = q("SELECT number, company_id FROM invoices WHERE id = ?", (iid,), one=True) if iid else None
    n = cur["number"] if cur else 10 ** 12
    base = f"SELECT i.id FROM invoices i JOIN customers c ON c.id = i.customer_id WHERE 1 = 1 {where}"
    co = cur["company_id"] if cur else (default_company_id() if len(companies()) > 1 else None)
    if co:
        base += " AND i.company_id = :navco"
        args["navco"] = co
    prev = q(base + " AND i.number < :n ORDER BY i.number DESC LIMIT 1", {**args, "n": n}, one=True)
    nxt = q(base + " AND i.number > :n ORDER BY i.number LIMIT 1", {**args, "n": n}, one=True) if cur else None
    return (prev["id"] if prev else None), (nxt["id"] if nxt else None)


def courier_list():
    """Couriers from Settings: one per line, 'Name | tracking link with {no}'."""
    out = []
    for line in (settings().get("couriers") or "").splitlines():
        name, _, link = line.partition("|")
        if name.strip():
            out.append({"name": name.strip(), "link": link.strip()})
    return out


def tracking_link(courier, number):
    if not courier or not number:
        return ""
    from urllib.parse import quote
    for c in courier_list():
        if c["name"].lower() == courier.lower() or c["name"].lower().startswith(courier.lower() + " "):
            return c["link"].replace("{no}", quote(number)) if c["link"] else ""
    return ""


def _render_invoice_form(inv, grid, f):
    s = settings()
    iid = inv["id"] if inv else None
    open_amt, cls, label = None, "", ""
    if inv:
        open_amt = open_items(inv["customer_id"])[0].get(iid, 0)
        cls, label = inv_status(inv, open_amt)
    prev_id, next_id = _nav_ids(iid)
    items_ = q("SELECT id, code, name, description, unit, price, kind, percent FROM items WHERE active = 1 "
               "ORDER BY kind != 'item', code")
    from .stock import levels
    lv = levels()
    item_data = [{"id": i["id"], "code": (i["code"] or i["name"]).upper(), "name": i["name"],
                  "stock": (f"{lv[i['id']]['pcs']:g} pcs" if i["id"] in lv else ""),
                  "desc": i["description"] or i["name"], "unit": i["unit"] or "", "kind": i["kind"],
                  "price": (f"{i['percent']:g}%" if i["kind"] == "discount" else
                            ("" if i["kind"] == "subtotal" else plain(i["price"])))} for i in items_]
    vias = [r["via"] for r in q("SELECT DISTINCT via FROM invoices WHERE via != '' ORDER BY via LIMIT 50")]
    msgs = [r["customer_message"] for r in q("""SELECT DISTINCT customer_message FROM invoices
                                               WHERE customer_message != '' ORDER BY customer_message LIMIT 30""")]
    sent = notify.last_sent("invoice", iid) if iid else None
    dcs = q("SELECT id, number FROM wh_docs WHERE invoice_id = ? AND type = 'DC' AND void = 0 ORDER BY id", (iid,)) if iid else []
    bid = (f.get("batch_id") if f else None) or (inv["batch_id"] if inv else None)
    batch = q("SELECT * FROM po_batches WHERE id = ?", (bid,), one=True) if bid else None
    return render_template("invoice_form.html", inv=inv, grid=grid, f=f, s=s, customers=scoped_customer_list(), dcs=dcs,
                           hist=invoice_history(inv) if inv else None,
                           couriers=courier_list(), batch=batch,
                           reps=q("SELECT * FROM reps WHERE active = 1 ORDER BY name"), item_data=item_data,
                           vias=vias, cust_msgs=msgs, open_amt=open_amt, cls=cls, label=label, prev_id=prev_id,
                           next_id=next_id, next_number=next_invoice_number(inv["company_id"] if inv else default_company_id()), sent=sent,
                           companies=companies(), inv_company=get_company(inv["company_id"]) if inv else None)


@bp.route("/invoices/<int:iid>/edit", methods=["GET", "POST"])
def invoice_form(iid=None):
    inv = q("SELECT * FROM invoices WHERE id = ?", (iid,), one=True) if iid else None
    if iid:
        if not inv:
            abort(404)
        get_customer(inv["customer_id"])
        if not can("invoices.edit") or inv["void"]:
            return redirect(url_for("sales.invoice_view", iid=iid, summary=1))

    if request.method == "POST":
        f = request.form
        try:
            cid = int(f.get("customer_id") or 0)
            c = get_customer(cid) if cid else None
            if not c:
                raise ValueError("Please choose a customer.")
            lines = _read_lines(f)
            if not any(l["kind"] != "subtotal" and (l["amount"] or l["description"]) for l in lines):
                raise ValueError("Add at least one item line.")
            d = parse_date(f.get("date")) or today()
            due = parse_date(f.get("due_date")) or (date.fromisoformat(d) + timedelta(days=c["terms_days"] or 0)).isoformat()
            rep_id = g.user["rep_id"] if is_rep() else (int(f["rep_id"]) if f.get("rep_id") else c["rep_id"])
            tax_rate = float((f.get("tax_rate") or "0").replace("%", "") or 0)
            co = get_company(c["company_id"]) or get_company(default_company_id())
            if not co["gst_registered"]:
                tax_rate = 0.0
            if inv and inv["company_id"] and inv["company_id"] != co["id"]:
                raise ValueError("This invoice belongs to another company. Choose a customer of the same company, "
                                 "or void it and make a new invoice.")
            branch_id = f.get("branch_id", type=int)
            if branch_id and not q("SELECT 1 FROM customer_branches WHERE id = ? AND customer_id = ?", (branch_id, cid)):
                raise ValueError("That branch belongs to another customer.")
            batch_id = f.get("batch_id", type=int)
            if batch_id and not q("SELECT 1 FROM po_batches WHERE id = ? AND customer_id = ? AND void = 0", (batch_id, cid)):
                batch_id = None
            if inv and not batch_id and "batch_id" not in f:
                batch_id = inv["batch_id"]
            number = None
            if not is_rep() and f.get("number", "").strip():
                number = int(f["number"].strip())
                clash = q("SELECT id FROM invoices WHERE number = ? AND company_id = ? AND id != ?", (number, co["id"], iid or -1), one=True)
                if clash:
                    raise ValueError(f"Invoice number {number} is already used for {co['name']}.")
        except ValueError as e:
            msg = str(e)
            if "invalid literal" in msg:
                msg = "Invoice number must be a whole number."
            flash(msg, "err")
            return _render_invoice_form(inv, _form_lines(f), f)
        subtotal = sum(l["amount"] for l in lines if l["kind"] != "subtotal")
        tax = pct_rnd(subtotal, tax_rate)
        total = subtotal + tax
        vals = dict(date=d, due_date=due, customer_id=cid, rep_id=rep_id, notes=f.get("notes", "").strip(),
                    tax_rate=tax_rate, subtotal=subtotal, discount=0, tax=tax, total=total, t=now(),
                    bilti_no=f.get("bilti_no", "").strip(), via=f.get("via", "").strip(),
                    ref_no=f.get("ref_no", "").strip(), ctn_count=f.get("ctn_count", "").strip(), po_no=f.get("po_no", "").strip(),
                    customer_message=f.get("customer_message", "").strip(), branch_id=branch_id, batch_id=batch_id,
                    courier=f.get("courier", "").strip(), tracking_no=f.get("tracking_no", "").strip(),
                    deliver_to=f.get("deliver_to", "").strip())
        if branch_id and not vals["deliver_to"]:
            b_ = q("SELECT * FROM customer_branches WHERE id = ?", (branch_id,), one=True)
            vals["deliver_to"] = "\n".join(v for v in (b_["name"], b_["address"], b_["phone"]) if v)
        if inv:
            x("""UPDATE invoices SET number=COALESCE(:number, number), date=:date, due_date=:due_date,
                 customer_id=:customer_id, rep_id=:rep_id, notes=:notes, tax_rate=:tax_rate, subtotal=:subtotal,
                 discount=:discount, tax=:tax, total=:total, bilti_no=:bilti_no, via=:via, ref_no=:ref_no,
                 ctn_count=:ctn_count, po_no=:po_no, customer_message=:customer_message, branch_id=:branch_id, batch_id=:batch_id,
                 courier=:courier, tracking_no=:tracking_no, deliver_to=:deliver_to, updated_at=:t WHERE id=:id""",
              {**vals, "number": number, "id": iid})
            x("DELETE FROM invoice_lines WHERE invoice_id = ?", (iid,))
        else:
            vals["number"] = number or next_invoice_number(co["id"])
            iid = x("""INSERT INTO invoices (number, date, due_date, customer_id, rep_id, notes, tax_rate, subtotal,
                       discount, tax, total, bilti_no, via, ref_no, ctn_count, po_no, customer_message, created_by, created_at, company_id,
                       branch_id, batch_id, courier, tracking_no, deliver_to)
                       VALUES (:number, :date, :due_date, :customer_id, :rep_id, :notes, :tax_rate, :subtotal,
                       :discount, :tax, :total, :bilti_no, :via, :ref_no, :ctn_count, :po_no, :customer_message, :uid, :t, :co,
                       :branch_id, :batch_id, :courier, :tracking_no, :deliver_to)""",
                    {**vals, "uid": g.user["id"], "co": co["id"]})
        for l in lines:
            x("""INSERT INTO invoice_lines (invoice_id, item_id, code, description, unit, ctn_qty, qty, rate, percent,
                 amount, kind, sort) VALUES (:iid, :item_id, :code, :description, :unit, :ctn_qty, :qty, :rate,
                 :percent, :amount, :kind, :sort)""", {**l, "iid": iid})
        commit()
        saved = q("SELECT number FROM invoices WHERE id = ?", (iid,), one=True)["number"]
        flash(f"Invoice {saved} {'updated' if inv else 'saved'}: {settings().get('currency', 'Rs')} {fmt(total)}.", "ok")
        if c["credit_limit"] and balance(cid) > c["credit_limit"]:
            flash(f"Warning: {c['name']} is now over their credit limit.", "err")
        if not inv and settings().get("auto_send_invoice") == "1" and (c["whatsapp"] or c["email"]):
            for ok, msg in notify.send_by(c, lambda: notify.send_invoice(iid), lambda: notify.email_invoice(iid)):
                flash(msg, "ok" if ok else "err")
        action = f.get("action", "save")
        if action in ("save_wa", "save_wa_img"):
            ok, msg = notify.send_invoice(iid, how="image" if action == "save_wa_img" else None)
            flash(msg, "ok" if ok else "err")
        if action == "save_email":
            ok, msg = notify.email_invoice(iid)
            flash(msg, "ok" if ok else "err")
        if action in ("save_print", "save_dc"):
            return redirect(url_for("sales.invoice_form", iid=iid, open="dc" if action == "save_dc" else "pdf"))
        if action == "save_new":
            return redirect(url_for("sales.invoice_new"))
        if action == "save_close":
            return redirect(url_for("sales.invoices"))
        if action == "save_pay":
            return redirect(url_for("sales.payment_new", customer_id=cid))
        if action == "save_po" and batch_id:
            return redirect(url_for("po.view", bid=batch_id))
        return redirect(url_for("sales.invoice_form", iid=iid) if can("invoices.edit") else url_for("sales.invoice_view", iid=iid, summary=1))

    if inv:
        return _render_invoice_form(inv, _grid_rows(iid), {})
    start = {"customer_id": request.args.get("customer_id", "")}
    bid = request.args.get("po", type=int)
    b = q("SELECT * FROM po_batches WHERE id = ? AND void = 0", (bid,), one=True) if bid else None
    if b:
        start.update(customer_id=str(b["customer_id"]), po_no=b["po_no"], batch_id=b["id"],
                     branch_id=request.args.get("branch", type=int) or "")
    return _render_invoice_form(None, [], start)


@bp.route("/invoices/new", methods=["GET", "POST"])
def invoice_new():
    return invoice_form()


@bp.route("/invoices/go/<direction>")
def invoice_go(direction):
    prev_id, next_id = _nav_ids(request.args.get("from", type=int))
    target = prev_id if direction == "prev" else next_id
    if not target:
        flash("No more invoices in that direction.", "info")
        return redirect(request.referrer or url_for("sales.invoices"))
    return redirect(url_for("sales.invoice_view", iid=target))


def get_invoice(iid):
    inv = q("""SELECT i.*, c.name AS customer, c.whatsapp, r.code AS rep_code, r.name AS rep_name
               FROM invoices i JOIN customers c ON c.id = i.customer_id LEFT JOIN reps r ON r.id = i.rep_id
               WHERE i.id = ?""", (iid,), one=True)
    if not inv:
        abort(404)
    get_customer(inv["customer_id"])
    return inv


def invoice_history(inv):
    """Every payment on one invoice in date order, with the balance left after each (for part payments)."""
    iid = inv["id"]
    applied = apply_payments(inv["customer_id"])["applied"]
    pays = {r["id"]: r for r in q("""SELECT p.id, p.number, p.date, p.method, p.reference, p.created_at, u.name AS by_name
                                      FROM payments p LEFT JOIN users u ON u.id = p.created_by WHERE p.customer_id = ?""",
                                   (inv["customer_id"],))}
    rows = sorted(({"p": pays[pid_], "amount": a} for pid_, lst in applied.items() for i_, a in lst if i_ == iid and pid_ in pays),
                  key=lambda r: (r["p"]["date"], r["p"]["number"]))
    left = inv["total"]
    for r in rows:
        left -= r["amount"]
        r["left"] = left
    o = open_items(inv["customer_id"])[0].get(iid, 0) if not inv["void"] else 0
    other = left - o if not inv["void"] else 0  # returns / credit notes / opening credit used on this bill
    return {"rows": rows, "paid": inv["total"] - left, "credit": other if other > 0 else 0, "open": o}


@bp.route("/invoices/<int:iid>")
def invoice_view(iid):
    inv = get_invoice(iid)
    if can("invoices.edit") and not inv["void"] and not request.args.get("summary"):
        return redirect(url_for("sales.invoice_form", iid=iid))
    lines = q("SELECT * FROM invoice_lines WHERE invoice_id = ? ORDER BY sort", (iid,))
    o = open_items(inv["customer_id"])[0].get(iid, 0)
    cls, label = inv_status(inv, o)
    msgs = q("SELECT * FROM messages WHERE kind = 'invoice' AND ref_id = ? ORDER BY id DESC", (iid,))
    prev_id, next_id = _nav_ids(iid)
    hist = invoice_history(inv)
    paid_by = hist["rows"]
    dcs = q("SELECT id, number, date FROM wh_docs WHERE invoice_id = ? AND type = 'DC' AND void = 0 ORDER BY id", (iid,))
    inv_email = q("SELECT email FROM customers WHERE id = ?", (inv["customer_id"],), one=True)["email"]
    return render_template("invoice_view.html", inv_email=inv_email, paid_by=paid_by, hist=hist, dcs=dcs, inv=inv, lines=lines, open=o, cls=cls, label=label, msgs=msgs,
                           wa=notify.normalize(inv["whatsapp"]) or "", prev_id=prev_id, next_id=next_id,
                           qty_text=qty_text, price_text=price_text)


@bp.route("/invoices/<int:iid>.pdf")
def invoice_pdf(iid):
    get_invoice(iid)
    return send_file(notify.invoice_pdf(iid), mimetype="application/pdf")


IMPORT_COLS = {
    "code": ("item code", "code", "item", "sku", "product code", "item no", "item #"),
    "qty": ("qty", "qty.", "quantity", "pcs"),
    "price": ("price", "price each", "rate", "unit price", "sale price"),
    "description": ("description", "desc", "item name", "name", "product"),
    "unit": ("u/m", "um", "unit", "uom"),
    "ctn_qty": ("ctn qty", "ctn", "carton", "cartons", "ctn qty."),
}


def _read_sheet(fs):
    name = (fs.filename or "").lower()
    data = fs.read()
    if name.endswith(".csv"):
        import csv
        import io
        text = data.decode("utf-8-sig", errors="replace")
        return [r for r in csv.reader(io.StringIO(text))]
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return [["" if v is None else v for v in r] for r in wb.active.iter_rows(values_only=True)]
    raise ValueError("Please save the file as Excel (.xlsx) or CSV first.")


@bp.route("/invoices/import-lines", methods=["POST"])
def invoice_import_lines():
    """Reads item lines from an Excel/CSV file for the invoice screen. Nothing is saved until the invoice is saved."""
    from .ops import item_data
    fs = request.files.get("file")
    if not fs:
        return jsonify(error="Choose a file."), 400
    try:
        rows = _read_sheet(fs)
    except Exception as e:  # noqa: BLE001 - show the user what went wrong
        return jsonify(error=str(e) if isinstance(e, ValueError) else "Couldn't read that file. Save it as .xlsx or .csv and try again."), 400
    rows = [[str(v).strip() if not isinstance(v, (int, float)) else v for v in r] for r in rows if any(str(v).strip() for v in r)]
    cols, start = {}, 0
    for n, r in enumerate(rows[:10]):
        found = {}
        for k, names in IMPORT_COLS.items():
            for i, v in enumerate(r):
                if str(v).strip().lower() in names and k not in found and i not in found.values():
                    found[k] = i
        if "code" in found or ("description" in found and "qty" in found):
            cols, start = found, n + 1
            break
    if not cols:
        cols = {"code": 0, "qty": 1, "price": 2}
    items = {i["code"]: i for i in item_data()}
    out, unknown = [], []

    def cell(r, k):
        i = cols.get(k)
        return r[i] if i is not None and i < len(r) else ""

    for r in rows[start:]:
        code = str(cell(r, "code")).strip().upper()
        if isinstance(cell(r, "code"), float) and float(cell(r, "code")).is_integer():
            code = str(int(cell(r, "code")))
        qty, price = cell(r, "qty"), cell(r, "price")
        desc = str(cell(r, "description")).strip()
        if not code and not desc:
            continue
        it = items.get(code)
        if code and not it:
            unknown.append(code)
        fmtn = lambda v: (f"{v:g}" if isinstance(v, float) else str(v)).strip()  # noqa: E731
        out.append({"code": code, "qty": fmtn(qty), "unit": str(cell(r, "unit")) or (it["unit"] if it else ""),
                    "ctn_qty": fmtn(cell(r, "ctn_qty")), "description": desc or (it["desc"] if it else ""),
                    "price": fmtn(price) if str(price).strip() != "" else (it["price"] if it else ""),
                    "item_id": it["id"] if it else "", "kind": "item"})
    if not out:
        return jsonify(error="No item lines found. The file needs columns like: Item code, Qty, Price."), 400
    return jsonify(rows=out, unknown=unknown[:20])


@bp.route("/invoices/import-template")
def invoice_import_template():
    from .export import table_response
    sample = q("SELECT code, description, unit, price FROM items WHERE active = 1 AND kind = 'item' ORDER BY code LIMIT 3")
    rows = [[i["code"], 10, i["unit"] or "", "", i["price"], i["description"] or ""] for i in sample] or [["ITEM1", 10, "Box", "", 25000, ""]]
    return table_response("xlsx", "Invoice lines import", ["Item code", "Qty", "U/M", "CTN qty", "Price", "Description"], rows, {4})


@bp.route("/invoices/<int:iid>/dc.pdf")
def invoice_dc_pdf(iid):
    get_invoice(iid)
    return send_file(notify.dc_pdf(iid), mimetype="application/pdf")


@bp.route("/invoices/<int:iid>/delivered", methods=["POST"])
def invoice_delivered(iid):
    inv = get_invoice(iid)
    if inv["delivered_on"]:
        x("UPDATE invoices SET delivered_on = '', delivered_by = '' WHERE id = ?", (iid,))
        flash(f"Invoice {inv['number']} marked as NOT delivered.", "ok")
    else:
        x("UPDATE invoices SET delivered_on = ?, delivered_by = ? WHERE id = ?",
          (parse_date(request.form.get("date")) or today(), g.user["name"], iid))
        flash(f"Invoice {inv['number']} marked as delivered.", "ok")
    commit()
    back = request.form.get("back", "")
    return redirect(back if back.startswith("/") and not back.startswith("//") else url_for("sales.invoice_form", iid=iid))


@bp.route("/invoices/<int:iid>/email", methods=["POST"])
def invoice_email(iid):
    inv = get_invoice(iid)
    if inv["void"]:
        abort(400, "A void invoice can't be sent.")
    ok, msg = notify.email_invoice(iid, request.form.get("email"))
    flash(msg, "ok" if ok else "err")
    back = request.form.get("back", "")
    return redirect(back if back.startswith("/") and not back.startswith("//") else url_for("sales.invoice_view", iid=iid))


@bp.route("/invoices/<int:iid>/send", methods=["POST"])
def invoice_send(iid):
    inv = get_invoice(iid)
    if inv["void"]:
        abort(400, "A void invoice can't be sent.")
    ok, msg = notify.send_invoice(iid, request.form.get("phone"), _fmt())
    flash(msg, "ok" if ok else "err")
    return redirect(request.form.get("back") or url_for("sales.invoice_view", iid=iid))


@bp.route("/invoices/<int:iid>/void", methods=["POST"])
@roles("admin", "accounts")
def invoice_void(iid):
    from .approvals import void_or_request
    get_invoice(iid)
    flash(*void_or_request("invoice", iid, request.form.get("reason", "").strip()))
    return redirect(url_for("sales.invoice_view", iid=iid, summary=1))


# ---- payments ----------------------------------------------------------------------------
@bp.route("/payments")
def payments():
    where, args = rep_filter()
    a = request.args
    if a.get("q"):
        where += " AND (c.name LIKE :s OR p.reference LIKE :s OR CAST(p.number AS TEXT) = :n)"
        args.update(s=f"%{a['q'].strip()}%", n=a["q"].strip())
    if a.get("method") in METHODS:
        where += " AND p.method = :m"
        args["m"] = a["method"]
    if parse_date(a.get("from")):
        where += " AND p.date >= :from"
        args["from"] = a["from"]
    if parse_date(a.get("to")):
        where += " AND p.date <= :to"
        args["to"] = a["to"]
    rows = q(f"""SELECT p.*, c.name AS customer FROM payments p JOIN customers c ON c.id = p.customer_id
                 WHERE 1 = 1 {where} ORDER BY p.date DESC, p.number DESC LIMIT 500""", args)
    if a.get("export") in ("xlsx", "pdf"):
        data = [[p["number"], nice_date(p["date"]), p["customer"], p["method"], p["reference"] or "",
                 "Bill-wise" if p["apply_mode"] == "manual" else "Auto", 0 if p["void"] else p["amount"],
                 "VOID" if p["void"] else ""] for p in rows]
        return table_response(a["export"], "Payments received", ["Receipt", "Date", "Customer", "Method", "Reference",
                              "Applied", "Amount", ""], data, {6}, ["Total", "", "", "", "", "", sum(d[6] for d in data), ""])
    return render_template("payments.html", rows=rows, methods=METHODS,
                           total=sum(p["amount"] for p in rows if not p["void"]))


TX_TYPES = {"invoice": "Invoice", "payment": "Payment", "credit": "Sales return"}


def _tx_rows(a, limit=1000):
    """Invoices, payments and returns in one list, with one set of filters (the 'All transactions' screen)."""
    where, args = rep_filter()
    parts = []
    types = [t for t in a.getlist("type") if t in TX_TYPES] or list(TX_TYPES)
    common = where
    if parse_date(a.get("from")):
        args["from"] = a["from"]
    if parse_date(a.get("to")):
        args["to"] = a["to"]
    dt = lambda col: (f" AND {col} >= :from" if "from" in args else "") + (f" AND {col} <= :to" if "to" in args else "")
    if "invoice" in types:
        parts.append(f"""SELECT 'invoice' AS t, i.id, i.number, i.date, c.id AS cid, c.name AS customer, c.code AS ccode, i.total AS amount, i.void,
                         TRIM(COALESCE(i.po_no, '') || ' ' || COALESCE(i.ref_no, '')) AS ref, '' AS method, i.created_at
                         FROM invoices i JOIN customers c ON c.id = i.customer_id WHERE 1 = 1 {common} {dt('i.date')}""")
    if "payment" in types:
        parts.append(f"""SELECT 'payment' AS t, p.id, p.number, p.date, c.id AS cid, c.name AS customer, c.code AS ccode, p.amount AS amount, p.void,
                         COALESCE(p.reference, '') AS ref, p.method AS method, p.created_at
                         FROM payments p JOIN customers c ON c.id = p.customer_id WHERE 1 = 1 {common} {dt('p.date')}""")
    if "credit" in types:
        parts.append(f"""SELECT 'credit' AS t, n.id, n.number, n.date, c.id AS cid, c.name AS customer, c.code AS ccode, n.total AS amount, n.void,
                         COALESCE(n.invoice_ref, '') AS ref, '' AS method, n.created_at
                         FROM credit_notes n JOIN customers c ON c.id = n.customer_id WHERE 1 = 1 {common} {dt('n.date')}""")
    sql = " UNION ALL ".join(parts)
    outer = []
    s_ = (a.get("q") or "").strip()
    if s_:
        args["qs"], args["qn"] = f"%{s_}%", s_
        outer.append("(customer LIKE :qs OR ccode LIKE :qs OR ref LIKE :qs OR CAST(number AS TEXT) = :qn OR method LIKE :qs)")
    amt = (a.get("amount") or "").replace(",", "").strip()
    if amt:
        try:
            args["amt"] = to_paisa(amt)
            outer.append("amount = :amt")
        except Exception:
            pass
    if a.get("min"):
        args["amin"] = to_paisa(a["min"])
        outer.append("amount >= :amin")
    if a.get("max"):
        args["amax"] = to_paisa(a["max"])
        outer.append("amount <= :amax")
    if a.get("status") == "void":
        outer.append("void = 1")
    elif a.get("status") != "all":
        outer.append("void = 0")
    order = {"date_asc": "date, number", "amount_desc": "amount DESC", "amount_asc": "amount", "customer": "customer COLLATE NOCASE, date DESC"
             }.get(a.get("sort"), "date DESC, created_at DESC")
    return q(f"SELECT * FROM ({sql}) WHERE 1 = 1 {(' AND ' + ' AND '.join(outer)) if outer else ''} ORDER BY {order} LIMIT {int(limit)}", args)


@bp.route("/transactions")
def transactions():
    a = request.args
    allowed = [t for t in TX_TYPES if can({"invoice": "invoices.view", "payment": "payments.view", "credit": "returns.view"}[t])]
    rows = [r for r in _tx_rows(a) if r["t"] in allowed]
    if a.get("export") in ("xlsx", "pdf"):
        data = [[TX_TYPES[r["t"]], r["number"], nice_date(r["date"]), r["customer"], r["ref"], r["method"],
                 0 if r["void"] else (r["amount"] if r["t"] == "invoice" else -r["amount"]), "VOID" if r["void"] else ""] for r in rows]
        return table_response(a["export"], "All transactions", ["Type", "No.", "Date", "Customer", "Reference", "Method", "Amount", ""],
                              data, {6}, None)
    tot = {t: sum(r["amount"] for r in rows if r["t"] == t and not r["void"]) for t in TX_TYPES}
    filtered = any(a.get(k) for k in ("q", "amount", "min", "max", "from", "to", "status", "sort")) or bool(a.getlist("type"))
    return render_template("transactions.html", rows=rows, types=TX_TYPES, allowed=allowed, tot=tot, filtered=filtered,
                           sel_types=a.getlist("type"))


@bp.route("/search")
def search():
    """The big search box: customers, invoices, payments, POs and items in one place."""
    s_ = request.args.get("q", "").strip()
    res = {}
    if s_:
        like = f"%{s_}%"
        where, args = rep_filter()
        if can("customers.view"):
            res["customers"] = q(f"""SELECT c.* FROM customers c WHERE 1 = 1 {where} AND (c.name LIKE :s OR c.code LIKE :s OR c.phone LIKE :s
                                    OR c.whatsapp LIKE :s OR c.contact LIKE :s OR c.city LIKE :s OR c.address LIKE :s) ORDER BY c.name LIMIT 30""",
                                 {**args, "s": like})
        tx = _tx_rows(request.args.__class__({"q": s_, "status": "all"}), 60)
        res["tx"] = [r for r in tx if can({"invoice": "invoices.view", "payment": "payments.view", "credit": "returns.view"}[r["t"]])]
        num = s_.replace(",", "")
        try:
            v = to_paisa(num)
            if v:
                res["by_amount"] = [r for r in _tx_rows(request.args.__class__({"amount": num, "status": "all"}), 30)
                                    if can({"invoice": "invoices.view", "payment": "payments.view", "credit": "returns.view"}[r["t"]])]
        except Exception:
            pass
        if can("invoices.view"):
            res["pos"] = q(f"""SELECT b.*, c.name AS customer FROM po_batches b JOIN customers c ON c.id = b.customer_id
                               WHERE b.void = 0 {where} AND (b.po_no LIKE :s OR c.name LIKE :s) ORDER BY b.date DESC LIMIT 20""", {**args, "s": like})
        if can("items.view"):
            res["items"] = q("SELECT * FROM items WHERE active = 1 AND kind = 'item' AND (code LIKE ? OR name LIKE ? OR description LIKE ?) "
                             "ORDER BY code LIMIT 20", (like, like, like))
    return render_template("search.html", s=s_, res=res, types=TX_TYPES)


def _read_allocations(f):
    out = []
    for inv, amt in zip(f.getlist("alloc_invoice"), f.getlist("alloc_amount")):
        a = to_paisa(amt)
        if a > 0:
            out.append((None if inv == "opening" else int(inv), a))
    return out


def _payment_form(p=None):
    """New payment (p=None) or edit an existing one."""
    customers_ = scoped_customer_list()
    s = settings()
    pid = p["id"] if p else None
    if request.method == "POST":
        f = request.form
        try:
            cid = int(f.get("customer_id") or 0)
            c = get_customer(cid) if cid else None
            if not c:
                raise ValueError("Please choose a customer.")
            mode = "manual" if f.get("apply_mode") == "manual" else "auto"
            allocs = _read_allocations(f) if mode == "manual" else []
            amount = to_paisa(f.get("amount"))
            if mode == "manual" and not amount:
                amount = sum(a for _, a in allocs)
            if amount <= 0:
                raise ValueError("Payment amount must be more than zero.")
            if sum(a for _, a in allocs) > amount:
                raise ValueError("The amounts applied to bills are more than the payment amount.")
            open_, opening_left, _ = open_items(cid, exclude_payment=pid)
            for inv_id, a in allocs:
                room = opening_left if inv_id is None else open_.get(inv_id)
                if room is None:
                    raise ValueError("One of the selected bills doesn't belong to this customer.")
                if a > room:
                    num = "opening balance" if inv_id is None else "invoice " + str(
                        q("SELECT number FROM invoices WHERE id = ?", (inv_id,), one=True)["number"])
                    raise ValueError(f"Payment on {num} is more than its amount due ({fmt(room)}).")
        except ValueError as e:
            flash(str(e), "err")
            return render_template("payment_form.html", customers=customers_, methods=METHODS, f=f, p=p,
                                   last_id=_pay_nav(pid)[0], next_id=_pay_nav(pid)[1] if pid else None,
                                   posted_allocs=[{"invoice": i, "amount": a} for i, a in
                                                  zip(f.getlist("alloc_invoice"), f.getlist("alloc_amount"))])
        vals = (parse_date(f.get("date")) or today(), cid, amount,
                f.get("method") if f.get("method") in METHODS else "Other", f.get("reference", "").strip(),
                f.get("notes", "").strip(), mode)
        if p:
            x("""UPDATE payments SET date=?, customer_id=?, amount=?, method=?, reference=?, notes=?, apply_mode=?,
                 updated_at=? WHERE id=?""", vals + (now(), pid))
            x("DELETE FROM payment_allocations WHERE payment_id = ?", (pid,))
        else:
            pid = x("""INSERT INTO payments (number, date, customer_id, amount, method, reference, notes, apply_mode,
                       created_by, created_at) VALUES (?,?,?,?,?,?,?,?,?,?)""",
                    (next_number("payments", "next_payment_number"),) + vals + (g.user["id"], now()))
        for inv_id, a in allocs:
            x("INSERT INTO payment_allocations (payment_id, invoice_id, amount) VALUES (?,?,?)", (pid, inv_id, a))
        x("UPDATE payments SET company_id = (SELECT company_id FROM customers c WHERE c.id = payments.customer_id) WHERE id = ?", (pid,))
        commit()
        left = amount - sum(a for _, a in allocs)
        how = "applied oldest bills first" if mode == "auto" else (
            f"applied to {len(allocs)} bill(s)" + (f", {fmt(left)} kept as advance" if left else ""))
        flash(f"Payment of {s.get('currency', 'Rs')} {fmt(amount)} from {c['name']} {'updated' if p else 'saved'} ({how}).", "ok")
        if not p and s.get("auto_send_receipt") == "1" and (c["whatsapp"] or c["email"]):
            for ok, msg in notify.send_by(c, lambda: notify.send_receipt(pid), lambda: notify.email_receipt(pid)):
                flash(msg, "ok" if ok else "err")
        action = f.get("action", "save")
        if action == "save_new":
            return redirect(url_for("sales.payment_new"))
        if action == "save_close":
            return redirect(url_for("sales.payments"))
        return redirect(url_for("sales.payment_view", pid=pid))

    if p:
        f = {"customer_id": p["customer_id"], "amount": plain(p["amount"]), "date": p["date"], "method": p["method"],
             "reference": p["reference"], "notes": p["notes"], "apply_mode": p["apply_mode"]}
        allocs = [{"invoice": a["invoice_id"] if a["invoice_id"] else "opening", "amount": plain(a["amount"])}
                  for a in q("SELECT * FROM payment_allocations WHERE payment_id = ?", (pid,))]
    else:
        cid = request.args.get("customer_id", "")
        f = {"customer_id": cid, "method": "Cash"}
        allocs = []
        bid = request.args.get("po", type=int)
        b = q("SELECT * FROM po_batches WHERE id = ? AND void = 0", (bid,), one=True) if bid else None
        if b:  # head office paying a whole PO: tick every open branch invoice of it
            open_ = open_items(b["customer_id"])[0]
            for i in q("SELECT id FROM invoices WHERE batch_id = ? AND void = 0 ORDER BY date, number", (bid,)):
                if open_.get(i["id"], 0) > 0:
                    allocs.append({"invoice": i["id"], "amount": plain(open_[i["id"]])})
            total = sum(open_.get(i["id"], 0) for i in q("SELECT id FROM invoices WHERE batch_id = ? AND void = 0", (bid,)))
            f.update(customer_id=b["customer_id"], apply_mode="manual", amount=plain(total) if total else "",
                     reference=f"PO {b['po_no']}", notes=f"Payment for PO {b['po_no']}" +
                     (f" (combined invoice {b['combined_no']})" if b["combined_no"] else ""))
    prev_id, next_id = _pay_nav(pid)
    return render_template("payment_form.html", customers=customers_, methods=METHODS, f=f, p=p,
                           last_id=prev_id, next_id=next_id if pid else None, posted_allocs=allocs)


@bp.route("/payments/new", methods=["GET", "POST"])
def payment_new():
    return _payment_form()


@bp.route("/payments/<int:pid>/edit", methods=["GET", "POST"])
@roles("admin", "accounts")
def payment_edit(pid):
    p = get_payment(pid)
    if p["void"]:
        flash("A void payment can't be edited.", "err")
        return redirect(url_for("sales.payment_view", pid=pid))
    return _payment_form(p)


def _pay_nav(pid):
    where, args = rep_filter()
    cur = q("SELECT number FROM payments WHERE id = ?", (pid,), one=True) if pid else None
    n = cur["number"] if cur else 10 ** 12
    base = f"SELECT p.id FROM payments p JOIN customers c ON c.id = p.customer_id WHERE 1 = 1 {where}"
    prev = q(base + " AND p.number < :n ORDER BY p.number DESC LIMIT 1", {**args, "n": n}, one=True)
    nxt = q(base + " AND p.number > :n ORDER BY p.number LIMIT 1", {**args, "n": n}, one=True) if cur else None
    return (prev["id"] if prev else None), (nxt["id"] if nxt else None)


@bp.route("/payments/go/<direction>")
def payment_go(direction):
    prev_id, next_id = _pay_nav(request.args.get("from", type=int))
    target = prev_id if direction == "prev" else next_id
    if not target:
        flash("No more payments in that direction.", "info")
        return redirect(request.referrer or url_for("sales.payments"))
    return redirect(url_for("sales.payment_view", pid=target))


def get_payment(pid):
    p = q("""SELECT p.*, c.name AS customer, c.whatsapp FROM payments p JOIN customers c ON c.id = p.customer_id
             WHERE p.id = ?""", (pid,), one=True)
    if not p:
        abort(404)
    get_customer(p["customer_id"])
    return p


@bp.route("/payments/<int:pid>")
def payment_view(pid):
    p = get_payment(pid)
    msgs = q("SELECT * FROM messages WHERE kind = 'receipt' AND ref_id = ? ORDER BY id DESC", (pid,))
    prev_id, next_id = _pay_nav(pid)
    applied = apply_payments(p["customer_id"])["applied"].get(pid, [])
    nums = {r["id"]: r for r in q("SELECT id, number, date, total FROM invoices WHERE customer_id = ?", (p["customer_id"],))}
    went = [{"inv": nums.get(i), "amount": a} for i, a in applied]
    c_email = q("SELECT email FROM customers WHERE id = ?", (p["customer_id"],), one=True)["email"]
    return render_template("payment_view.html", c_email=c_email, p=p, bal=balance(p["customer_id"]), msgs=msgs, went=went,
                           unapplied=p["amount"] - sum(a for _, a in applied),
                           wa=notify.normalize(p["whatsapp"]) or "", prev_id=prev_id, next_id=next_id)


@bp.route("/payments/<int:pid>.pdf")
def payment_pdf(pid):
    get_payment(pid)
    return send_file(notify.receipt_pdf(pid), mimetype="application/pdf")


@bp.route("/payments/<int:pid>/send", methods=["POST"])
def payment_send(pid):
    p = get_payment(pid)
    if p["void"]:
        abort(400, "A void payment can't be sent.")
    ok, msg = notify.send_receipt(pid, request.form.get("phone"), _fmt())
    flash(msg, "ok" if ok else "err")
    return redirect(url_for("sales.payment_view", pid=pid))


@bp.route("/payments/<int:pid>/email", methods=["POST"])
def payment_email(pid):
    p = get_payment(pid)
    if p["void"]:
        abort(400, "A void payment can't be sent.")
    ok, msg = notify.email_receipt(pid, request.form.get("email"))
    flash(msg, "ok" if ok else "err")
    return redirect(url_for("sales.payment_view", pid=pid))


@bp.route("/payments/<int:pid>/void", methods=["POST"])
@roles("admin", "accounts")
def payment_void(pid):
    from .approvals import void_or_request
    get_payment(pid)
    flash(*void_or_request("payment", pid, request.form.get("reason", "").strip()))
    return redirect(url_for("sales.payment_view", pid=pid))
