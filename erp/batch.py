"""Batch billing: one customer PO delivered to many branches.

Each branch gets its own invoice (bill to head office, ship to the branch) with its own DC and courier tracking.
When the deliveries are done, one combined invoice lists every branch invoice for head office. The branch invoices
carry the amounts in the customer's account; the combined invoice is the summary, so nothing is counted twice.
Head office then pays once and the payment clears all branch invoices of the PO.
"""
import io

from flask import Blueprint, abort, flash, g, redirect, render_template, request, send_file, url_for

from . import notify
from .auth import can, rep_filter
from .db import commit, fmt, nice_date, now, parse_date, plain, q, settings, today, x
from .ledger import open_items

bp = Blueprint("po", __name__, url_prefix="/po")


# ---- helpers ---------------------------------------------------------------------------------
def get_batch(bid):
    from .sales import get_customer
    b = q("""SELECT b.*, c.name AS customer, c.email, c.whatsapp, c.company_id AS cust_co FROM po_batches b
             JOIN customers c ON c.id = b.customer_id WHERE b.id = ?""", (bid,), one=True)
    if not b:
        abort(404)
    get_customer(b["customer_id"])
    return b


def batch_rows(b):
    """Every active branch of the customer with its invoice(s) on this PO, plus invoices without a branch."""
    from .sales import tracking_link
    invs = q("""SELECT i.*, (SELECT COUNT(*) FROM wh_docs d WHERE d.invoice_id = i.id AND d.void = 0 AND d.type = 'DC') AS wh_dc
                FROM invoices i WHERE i.batch_id = ? AND i.void = 0 ORDER BY i.number""", (b["id"],))
    open_ = open_items(b["customer_id"])[0]
    branches = q("SELECT * FROM customer_branches WHERE customer_id = ? ORDER BY active DESC, name COLLATE NOCASE", (b["customer_id"],))
    by_branch = {}
    for i in invs:
        by_branch.setdefault(i["branch_id"], []).append(i)
    rows = []
    for br in branches:
        mine = by_branch.pop(br["id"], [])
        if not mine and not br["active"]:
            continue
        rows.append({"branch": br, "invoices": [dict(i, open=open_.get(i["id"], 0), link=tracking_link(i["courier"], i["tracking_no"]))
                                                for i in mine]})
    for bid_, mine in by_branch.items():  # invoices with no branch, or a branch that was deleted
        rows.append({"branch": None, "invoices": [dict(i, open=open_.get(i["id"], 0), link=tracking_link(i["courier"], i["tracking_no"]))
                                                  for i in mine]})
    all_inv = [i for r in rows for i in r["invoices"]]
    t = {"subtotal": sum(i["subtotal"] for i in all_inv), "tax": sum(i["tax"] for i in all_inv),
         "total": sum(i["total"] for i in all_inv), "open": sum(i["open"] for i in all_inv),
         "count": len(all_inv), "delivered": sum(1 for i in all_inv if i["delivered_on"]),
         "shipped": sum(1 for i in all_inv if i["tracking_no"] or i["delivered_on"]),
         "branches": len([r for r in rows if r["branch"]]), "billed_branches": len([r for r in rows if r["branch"] and r["invoices"]])}
    t["paid"] = t["total"] - t["open"]
    return rows, all_inv, t


def status_of(b, t):
    if b["void"]:
        return "void", "Void"
    if t["count"] and t["open"] <= 0:
        return "paid", "Paid"
    if b["combined_no"]:
        return "sent", "Billed to head office"
    if t["count"] and t["delivered"] == t["count"]:
        return "open", "All delivered - make combined invoice"
    if t["count"]:
        return "open", f"{t['delivered']} of {t['count']} delivered"
    return "", "No branch invoices yet"


def next_combined_no(company_id):
    r = q("SELECT MAX(combined_no) AS n FROM po_batches WHERE company_id IS ?", (company_id,), one=True)
    return (r["n"] or 0) + 1


def combined_pdf(bid, draft=False):
    from .companies import co_images, co_settings
    from .pdfs import po_invoice
    b = q("SELECT * FROM po_batches WHERE id = ?", (bid,), one=True)
    c = q("SELECT * FROM customers WHERE id = ?", (b["customer_id"],), one=True)
    rows, all_inv, t = batch_rows(b)
    lines = {}
    if all_inv:
        ids = [i["id"] for i in all_inv]
        for l in q(f"SELECT * FROM invoice_lines WHERE invoice_id IN ({','.join('?' * len(ids))}) ORDER BY invoice_id, sort", ids):
            lines.setdefault(l["invoice_id"], []).append(l)
    s = co_settings(b["company_id"])
    s["_logo"] = co_images(b["company_id"]).get("logo")
    name = f"PO_{s.get('company_code', '')}_{notify._safe(b['po_no'])}.pdf"
    return po_invoice(notify._pdf_path(name), s, b, c, rows, t, lines, draft=draft)


# ---- screens ---------------------------------------------------------------------------------
@bp.route("/")
def index():
    where, args = rep_filter()
    st = request.args.get("status", "")
    rows = []
    for b in q(f"""SELECT b.*, c.name AS customer FROM po_batches b JOIN customers c ON c.id = b.customer_id
                   WHERE b.void = 0 {where} ORDER BY b.date DESC, b.id DESC LIMIT 500""", args):
        _, _, t = batch_rows(b)
        cls, label = status_of(b, t)
        if st == "open" and (b["combined_no"] or cls == "paid"):
            continue
        if st == "billed" and not (b["combined_no"] and cls != "paid"):
            continue
        if st == "paid" and cls != "paid":
            continue
        rows.append({"b": b, "t": t, "cls": cls, "label": label})
    return render_template("po/index.html", rows=rows, st=st)


@bp.route("/new", methods=["GET", "POST"])
def new():
    from .sales import get_customer, scoped_customer_list
    f = request.form if request.method == "POST" else {"customer_id": request.args.get("customer_id", ""), "date": today()}
    if request.method == "POST":
        cid = request.form.get("customer_id", type=int)
        c = get_customer(cid) if cid else None
        po_no = request.form.get("po_no", "").strip()
        if not c:
            flash("Choose the head-office customer.", "err")
        elif not po_no:
            flash("Enter the customer's PO number.", "err")
        elif q("SELECT 1 FROM po_batches WHERE customer_id = ? AND po_no = ? AND void = 0", (cid, po_no)):
            flash(f"PO {po_no} already exists for {c['name']}.", "err")
        else:
            bid = x("""INSERT INTO po_batches (company_id, customer_id, po_no, date, notes, created_by, created_at)
                       VALUES (?,?,?,?,?,?,?)""", (c["company_id"], cid, po_no, parse_date(request.form.get("date")) or today(),
                                                 request.form.get("notes", "").strip(), g.user["id"], now()))
            commit()
            flash(f"PO {po_no} created. Now make an invoice for each branch.", "ok")
            return redirect(url_for("po.view", bid=bid))
    return render_template("po/new.html", f=f, customers=scoped_customer_list())


@bp.route("/<int:bid>")
def view(bid):
    b = get_batch(bid)
    rows, all_inv, t = batch_rows(b)
    cls, label = status_of(b, t)
    from .sales import courier_list
    msgs = q("SELECT * FROM messages WHERE kind = 'po' AND ref_id = ? ORDER BY id DESC LIMIT 10", (bid,))
    loose = q("""SELECT id, number, date, total, branch_id FROM invoices WHERE customer_id = ? AND void = 0 AND batch_id IS NULL
                 ORDER BY date DESC, number DESC LIMIT 60""", (b["customer_id"],))
    return render_template("po/view.html", b=b, rows=rows, t=t, cls=cls, label=label, msgs=msgs, loose=loose,
                           couriers=courier_list())


@bp.route("/<int:bid>/edit", methods=["POST"])
def edit(bid):
    b = get_batch(bid)
    po_no = request.form.get("po_no", "").strip() or b["po_no"]
    if q("SELECT 1 FROM po_batches WHERE customer_id = ? AND po_no = ? AND void = 0 AND id != ?", (b["customer_id"], po_no, bid)):
        flash(f"PO {po_no} already exists for this customer.", "err")
        return redirect(url_for("po.view", bid=bid))
    x("UPDATE po_batches SET po_no = ?, date = ?, notes = ? WHERE id = ?",
      (po_no, parse_date(request.form.get("date")) or b["date"], request.form.get("notes", "").strip(), bid))
    if po_no != b["po_no"]:
        x("UPDATE invoices SET po_no = ? WHERE batch_id = ? AND (po_no = ? OR po_no = '')", (po_no, bid, b["po_no"]))
    commit()
    flash("PO updated.", "ok")
    return redirect(url_for("po.view", bid=bid))


@bp.route("/<int:bid>/tracking/<int:iid>", methods=["POST"])
def tracking(bid, iid):
    get_batch(bid)
    inv = q("SELECT * FROM invoices WHERE id = ? AND batch_id = ?", (iid, bid), one=True) or abort(404)
    f = request.form
    courier, trk = f.get("courier", "").strip(), f.get("tracking_no", "").strip()
    shipped = inv["shipped_on"] or (today() if trk else "")
    x("UPDATE invoices SET courier = ?, tracking_no = ?, shipped_on = ? WHERE id = ?", (courier, trk, shipped, iid))
    if f.get("delivered") == "1" and not inv["delivered_on"]:
        x("UPDATE invoices SET delivered_on = ?, delivered_by = ? WHERE id = ?", (today(), g.user["name"], iid))
    elif f.get("delivered") == "0" and inv["delivered_on"]:
        x("UPDATE invoices SET delivered_on = '', delivered_by = '' WHERE id = ?", (iid,))
    commit()
    flash(f"Invoice {inv['number']}: tracking saved.", "ok")
    return redirect(url_for("po.view", bid=bid) + f"#inv{iid}")


@bp.route("/<int:bid>/combine", methods=["POST"])
def combine(bid):
    b = get_batch(bid)
    _, all_inv, t = batch_rows(b)
    if not all_inv:
        flash("Make the branch invoices first.", "err")
    elif request.form.get("undo") == "1":
        x("UPDATE po_batches SET combined_no = NULL, combined_date = NULL WHERE id = ?", (bid,))
        commit()
        flash("Combined invoice cancelled. You can add more branches and make it again.", "ok")
    else:
        no = b["combined_no"] or next_combined_no(b["company_id"])
        x("UPDATE po_batches SET combined_no = ?, combined_date = ? WHERE id = ?",
          (no, parse_date(request.form.get("date")) or today(), bid))
        commit()
        left = t["count"] - t["delivered"]
        flash(f"Combined invoice {no} made for PO {b['po_no']}: {settings().get('currency', 'Rs')} {fmt(t['total'])}"
              + (f". Note: {left} branch invoice(s) are not marked delivered yet." if left else "."), "ok" if not left else "err")
    return redirect(url_for("po.view", bid=bid))


@bp.route("/<int:bid>/combined.pdf")
def pdf(bid):
    b = get_batch(bid)
    return send_file(combined_pdf(bid, draft=not b["combined_no"]), mimetype="application/pdf")


@bp.route("/<int:bid>/send", methods=["POST"])
def send(bid):
    b = get_batch(bid)
    if not b["combined_no"]:
        flash("Make the combined invoice first.", "err")
        return redirect(url_for("po.view", bid=bid))
    ch = request.form.get("channel", "whatsapp")
    c = q("SELECT * FROM customers WHERE id = ?", (b["customer_id"],), one=True)
    fns = (lambda: send_whatsapp(bid), lambda: send_email(bid, request.form.get("email") or None))
    results = ([fns[0]()] if ch == "whatsapp" else [fns[1]()] if ch == "email" else notify.send_by(c, *fns))
    for ok, msg in results:
        flash(msg, "ok" if ok else "err")
    return redirect(url_for("po.view", bid=bid))


def _summary(bid):
    from .companies import co_settings
    b = q("SELECT * FROM po_batches WHERE id = ?", (bid,), one=True)
    c = q("SELECT * FROM customers WHERE id = ?", (b["customer_id"],), one=True)
    _, _, t = batch_rows(b)
    return b, c, t, co_settings(b["company_id"])


def send_whatsapp(bid):
    b, c, t, s = _summary(bid)
    cur = s.get("currency", "Rs")
    params = [c["name"], f"{b['combined_no']} (PO {b['po_no']})", f"{cur} {fmt(t['open'])}", nice_date(b["combined_date"])]
    return notify._deliver("po", c, bid, None, settings().get("wa_invoice_template"), combined_pdf(bid), params,
                           f"Combined invoice for PO {b['po_no']}")


def send_email(bid, to=None):
    b, c, t, s = _summary(bid)
    cur = s.get("currency", "Rs")
    lines = [f"Please find attached our combined invoice no. {b['combined_no']} for your PO {b['po_no']}.",
             f"It covers {t['count']} branch invoice(s), total {cur} {fmt(t['total'])}"
             + (f", of which {cur} {fmt(t['open'])} is due." if t["paid"] else "."),
             "Each branch invoice and delivery challan is listed with its courier tracking number."]
    return notify._deliver_email("po", c, bid, to, f"Combined invoice {b['combined_no']} - PO {b['po_no']} - {s.get('company_name', '')}",
                                 lines, combined_pdf(bid), f"Combined invoice for PO {b['po_no']}")


@bp.route("/<int:bid>/void", methods=["POST"])
def void(bid):
    b = get_batch(bid)
    x("UPDATE invoices SET batch_id = NULL WHERE batch_id = ?", (bid,))
    x("UPDATE po_batches SET void = 1 WHERE id = ?", (bid,))
    commit()
    flash(f"PO {b['po_no']} removed. Its branch invoices are kept as normal invoices.", "ok")
    return redirect(url_for("po.index"))


@bp.route("/<int:bid>/pay")
def pay(bid):
    """Receive payment with every open branch invoice of this PO ticked."""
    b = get_batch(bid)
    return redirect(url_for("sales.payment_new", customer_id=b["customer_id"], po=bid))


@bp.route("/<int:bid>/attach", methods=["POST"])
def attach(bid):
    """Adds invoices already made (e.g. before the PO was entered) to this PO."""
    b = get_batch(bid)
    ids = [int(v) for v in request.form.getlist("iid") if v.isdigit()]
    for iid in ids:
        x("UPDATE invoices SET batch_id = ?, po_no = CASE WHEN po_no = '' THEN ? ELSE po_no END "
          "WHERE id = ? AND customer_id = ? AND void = 0 AND batch_id IS NULL", (bid, b["po_no"], iid, b["customer_id"]))
    commit()
    flash(f"{len(ids)} invoice(s) added to PO {b['po_no']}.", "ok")
    return redirect(url_for("po.view", bid=bid))


@bp.route("/<int:bid>/detach/<int:iid>", methods=["POST"])
def detach(bid, iid):
    get_batch(bid)
    x("UPDATE invoices SET batch_id = NULL WHERE id = ? AND batch_id = ?", (iid, bid))
    commit()
    flash("Invoice taken off this PO (the invoice itself is kept).", "ok")
    return redirect(url_for("po.view", bid=bid))


# ---- branches --------------------------------------------------------------------------------
@bp.route("/branches/<int:cid>", methods=["POST"])
def branch_save(cid):
    from .sales import get_customer
    get_customer(cid)
    f = request.form
    bid = f.get("id", type=int)
    name = f.get("name", "").strip()
    if not name:
        flash("Give the branch a name.", "err")
    else:
        vals = (name, f.get("code", "").strip(), f.get("address", "").strip(), f.get("contact", "").strip(),
                f.get("phone", "").strip(), f.get("email", "").strip())
        if bid:
            x("UPDATE customer_branches SET name=?, code=?, address=?, contact=?, phone=?, email=?, active=? WHERE id=? AND customer_id=?",
              vals + (0 if f.get("inactive") else 1, bid, cid))
        else:
            x("""INSERT INTO customer_branches (name, code, address, contact, phone, email, customer_id, created_at)
                 VALUES (?,?,?,?,?,?,?,?)""", vals + (cid, now()))
        commit()
        flash(f"Branch '{name}' saved.", "ok")
    return redirect(url_for("sales.customer_view", cid=cid, tab="branches"))


@bp.route("/branches/<int:cid>/import", methods=["POST"])
def branch_import(cid):
    """Excel / CSV with columns: Branch name, Code, Address, Contact, Phone, Email."""
    from .sales import get_customer
    get_customer(cid)
    fl = request.files.get("file")
    if not fl or not fl.filename:
        flash("Choose an Excel or CSV file.", "err")
        return redirect(url_for("sales.customer_view", cid=cid, tab="branches"))
    try:
        if fl.filename.lower().endswith(".csv"):
            import csv
            data = list(csv.reader(io.StringIO(fl.read().decode("utf-8-sig", "replace"))))
        else:
            from openpyxl import load_workbook
            ws = load_workbook(io.BytesIO(fl.read()), read_only=True, data_only=True).active
            data = [["" if v is None else str(v).strip() for v in r] for r in ws.iter_rows(values_only=True)]
    except Exception:
        flash("Could not read that file. Save it as .xlsx or .csv and try again.", "err")
        return redirect(url_for("sales.customer_view", cid=cid, tab="branches"))
    keys = {"name": ("branch", "name", "branch name"), "code": ("code", "branch code"), "address": ("address",),
            "contact": ("contact", "contact person"), "phone": ("phone", "mobile", "whatsapp", "cell"), "email": ("email", "e-mail")}
    head = [h.strip().lower() for h in (data[0] if data else [])]
    col = {k: next((i for i, h in enumerate(head) if h in names), None) for k, names in keys.items()}
    if col["name"] is None:
        col, start = {"name": 0, "code": 1, "address": 2, "contact": 3, "phone": 4, "email": 5}, 0
    else:
        start = 1
    have = {r["name"].lower() for r in q("SELECT name FROM customer_branches WHERE customer_id = ?", (cid,))}
    n = 0
    for r in data[start:]:
        get = lambda k: (r[col[k]].strip() if col.get(k) is not None and col[k] < len(r) and r[col[k]] else "")
        name = get("name")
        if not name or name.lower() in have:
            continue
        x("""INSERT INTO customer_branches (customer_id, name, code, address, contact, phone, email, created_at)
             VALUES (?,?,?,?,?,?,?,?)""", (cid, name, get("code"), get("address"), get("contact"), get("phone"), get("email"), now()))
        have.add(name.lower())
        n += 1
    commit()
    flash(f"{n} branch(es) added.", "ok")
    return redirect(url_for("sales.customer_view", cid=cid, tab="branches"))
